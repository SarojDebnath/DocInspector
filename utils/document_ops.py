from __future__ import annotations
from pathlib import Path
import os
from typing import Iterable, List
from fastapi import UploadFile
from langchain.schema import Document
from docling.document_converter import DocumentConverter
from openpyxl import load_workbook
from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader, TextLoader,UnstructuredExcelLoader
from logger import GLOBAL_LOGGER as log
from exception.custom_exception import DocumentPortalException
from langchain_excel_loader import StructuredExcelLoader
from openpyxl import load_workbook 
import pandas as pd

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt",".xlsx",".pptx"}

def load_xlsx_as_markdown(path: Path) -> list[Document]:
    """Structured Excel parsing: first 100 rows per sheet, one row per chunk,
    plus separate chunks for long note-like fields. Adds useful metadata."""
    docs: list[Document] = []
    try:
        xls = pd.ExcelFile(path)  # type: ignore
        workbook = Path(path).name
        for sheet_name in xls.sheet_names:
            try:
                df = xls.parse(sheet_name=sheet_name, dtype=str, nrows=100)  # type: ignore
            except Exception:
                continue
            if df is None or df.empty:
                continue
            df = df.fillna("")

            headers = list(df.columns)
            header_line = " | ".join(headers)

            # Row-wise chunks
            for idx, row in df.iterrows():  # type: ignore
                values = [str(row.get(h, "")).strip() for h in headers]
                kv_pairs = [f"{h}: {v}" for h, v in zip(headers, values) if v]
                if not kv_pairs:
                    continue
                content = (
                    f"Workbook: {workbook}\nSheet: {sheet_name}\nRow: {idx}\n"
                    f"Headers: {header_line}\n"
                    + "\n".join(kv_pairs)
                )
                meta: dict = {
                    "source": str(path),
                    "workbook": workbook,
                    "sheet": sheet_name,
                    "row": int(idx),
                }
                for key in ["lot", "id", "ref", "reference", "pos", "position"]:
                    for h, v in zip(headers, values):
                        if key in h.lower() and v:
                            meta[key] = v
                docs.append(Document(page_content=content, metadata=meta))

            # Long paragraph chunks for notes/description fields
            note_like_cols = [
                h for h in headers
                if any(k in h.lower() for k in ["note", "comment", "desc", "observation", "issue", "problem"])  # generic text columns
            ]
            if note_like_cols:
                for idx, row in df.iterrows():  # type: ignore
                    for h in note_like_cols:
                        val = str(row.get(h, "")).strip()
                        if len(val) >= 200:
                            docs.append(
                                Document(
                                    page_content=f"Workbook: {workbook}\nSheet: {sheet_name}\nRow: {idx}\nField: {h}\n\n{val}",
                                    metadata={
                                        "source": str(path),
                                        "workbook": workbook,
                                        "sheet": sheet_name,
                                        "row": int(idx),
                                        "field": h,
                                    },
                                )
                            )
    except Exception as e:
        log.error("Failed to parse Excel via pandas", file=str(path), error=str(e))
        # Fallback: single doc via docling below
        try:
            converter = DocumentConverter()
            result = converter.convert(str(path))
            text_md = result.document.export_to_markdown()
            if not isinstance(text_md, str):
                try:
                    text_md = "\n".join(map(str, text_md))  # type: ignore
                except Exception:
                    text_md = str(text_md)
            docs.append(Document(page_content=text_md, metadata={"source": str(path)}))
        except Exception:
            pass
    return docs

def load_documents(paths: Iterable[Path]) -> List[Document]:
    """Load docs using appropriate loader based on extension."""
    docs: List[Document] = []
    try:
        converter = DocumentConverter()
        for p in paths:
            ext = p.suffix.lower()
            if ext == ".pdf":
                loader = PyPDFLoader(str(p))
                docs.extend(loader.load())
            elif ext == ".docx":
                loader = Docx2txtLoader(str(p))
                docs.extend(loader.load())
            elif ext == ".txt":
                loader = TextLoader(str(p), encoding="utf-8")
                docs.extend(loader.load())
            elif ext == ".xlsx":
                try:
                    docs.extend(load_xlsx_as_markdown(p))
                except Exception as e:
                    log.error("Failed to load Excel", file=str(p), error=str(e))
                    # Fallback: docling single document
                    try:
                        result = converter.convert(str(p))
                        text_md = result.document.export_to_markdown()
                        if not isinstance(text_md, str):
                            try:
                                text_md = "\n".join(map(str, text_md))  # type: ignore
                            except Exception:
                                text_md = str(text_md)
                        docs.append(Document(page_content=text_md, metadata={"source": str(p)}))
                    except Exception:
                        continue

            elif ext == ".pptx":
                result = converter.convert(str(p))
                text_md = result.document.export_to_markdown()
                if not isinstance(text_md, str):
                    try:
                        text_md = "\n".join(map(str, text_md))  # type: ignore
                    except Exception:
                        text_md = str(text_md)
                lc_doc = Document(page_content=text_md, metadata={"source": str(p)})
                docs.append(lc_doc)
            else:
                log.warning("Unsupported extension skipped", path=str(p))
                continue
            
        log.info("Documents loaded", count=len(docs))
        return docs
    except Exception as e:
        log.error("Failed loading documents", error=str(e))
        raise DocumentPortalException("Error loading documents", e) from e

def concat_for_analysis(docs: List[Document]) -> str:
    parts = []
    for d in docs:
        src = d.metadata.get("source") or d.metadata.get("file_path") or "unknown"
        parts.append(f"\n--- SOURCE: {src} ---\n{d.page_content}")
    return "\n".join(parts)

def concat_for_comparison(ref_docs: List[Document], act_docs: List[Document]) -> str:
    left = concat_for_analysis(ref_docs)
    right = concat_for_analysis(act_docs)
    return f"<<REFERENCE_DOCUMENTS>>\n{left}\n\n<<ACTUAL_DOCUMENTS>>\n{right}"

# ---------- Helpers ----------
class FastAPIFileAdapter:
    """Adapt FastAPI UploadFile -> .name + .getbuffer() API"""
    def __init__(self, uf: UploadFile):
        self._uf = uf
        self.name = uf.filename
    def getbuffer(self) -> bytes:
        self._uf.file.seek(0)
        return self._uf.file.read()

def read_pdf_via_handler(handler, path: str) -> str:
    if hasattr(handler, "read_pdf"):
        return handler.read_pdf(path)  # type: ignore
    if hasattr(handler, "read_"):
        return handler.read_(path)  # type: ignore
    raise RuntimeError("DocHandler has neither read_pdf nor read_ method.")